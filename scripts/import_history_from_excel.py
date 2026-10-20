#!/usr/bin/env python3
"""Turn the manual status workbook into a clean, validated history seed.

    python scripts/import_history_from_excel.py data/source/Broga_Competitor_Sales_Status_13102026.xlsx

Reads Sheet1 (one block per project, dates across columns, status as multi-line text) and writes
  data/history-seed.json   one record per project x date, same shape as data/teduh-history.json
  data/history-issues.csv  every problem found, with the source cell, so you can fix it at the source

The Excel stays read-only. Fixes live in data/history-corrections.json and are logged on each record.
"""
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
ATTR_LABELS = {"land area": "land area", "bua": "bua", "spa price": "spa price",
               "nett price": "nett price", "rebate": "rebate", "remarks": "remarks"}
STATUS_RE = re.compile(r"(\d+(?:\.\d+)?)\s*%\s*(Signed|Booked)\s*\(\s*(\d+)\s*/\s*(\d+)\s*units?\s*\)", re.I)
PCT_TOLERANCE = 1.0  # percentage points; the sheet rounds to whole numbers


def parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and re.fullmatch(r"\s*\d{1,2}/\d{1,2}/\d{4}\s*", v):
        return datetime.strptime(v.strip(), "%d/%m/%Y").date()
    return None


def find_blocks(ws):
    blocks = []
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str) and a.strip().lower() not in ATTR_LABELS and a.strip().lower() != "sales status" \
                and parse_date(b):
            blocks.append(r)
    return blocks


def project_ids(wb):
    ids = {}
    if "Teduh Search Reference" in wb.sheetnames:
        for row in wb["Teduh Search Reference"].iter_rows(min_row=2, values_only=True):
            if row[0] and row[2]:
                ids[str(row[0]).strip()] = str(row[2]).splitlines()[0].strip()
    return ids


def load_corrections():
    p = DATA / "history-corrections.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {"corrections": [], "group_labels": {}}


def main(xlsx):
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Sheet1"]
    ids = project_ids(wb)
    cfg = load_corrections()
    issues = []

    def issue(sev, code, project, dt, cell, detail):
        issues.append({"severity": sev, "code": code, "project": project, "date": dt or "", "cell": cell, "detail": detail})

    blocks = find_blocks(ws)
    records, header_dates = [], defaultdict(set)
    for bi, top in enumerate(blocks):
        end = blocks[bi + 1] if bi + 1 < len(blocks) else ws.max_row + 1
        project = ws.cell(top, 1).value.strip()
        rows = {}
        for r in range(top, end):
            label = str(ws.cell(r, 1).value or "").strip().lower()
            if label in ATTR_LABELS or label == "sales status":
                rows[label] = r
        if "sales status" not in rows:
            issue("error", "NO_STATUS_ROW", project, "", f"A{top}", "Block has no 'Sales Status' row")
            continue
        for c in range(2, ws.max_column + 1):
            head = ws.cell(top, c).value
            d = parse_date(head)
            col = ws.cell(top, c).column_letter
            if d is None:
                if head not in (None, ""):
                    issue("info", "NON_DATE_HEADER", project, "", f"{col}{top}", f"Header '{head}' is text in a date column; ignored")
                continue
            header_dates[c].add((d.isoformat(), project))
            if isinstance(head, str):
                issue("info", "TEXT_DATE", project, d.isoformat(), f"{col}{top}", "Date stored as text, not an Excel date; parsed as dd/mm/yyyy")
            cell = ws.cell(rows["sales status"], c)
            text = str(cell.value).strip() if cell.value is not None else ""
            if not text:
                continue
            found = STATUS_RE.findall(text)
            comps = [{"label": k.title(), "units": int(u), "total": int(t), "reportedPct": float(p),
                      "raw": f"{p}% {k.title()} ({u}/{t} units)"} for p, k, u, t in found]
            if not comps and not re.search(r"yet to (sign|launch)", text, re.I):
                issue("error", "UNPARSEABLE_STATUS", project, d.isoformat(), cell.coordinate, f"Could not read: {text!r}")
            static = {}
            for key, r in rows.items():
                if key != "sales status":
                    v = ws.cell(r, c).value
                    if v not in (None, ""):
                        static[key] = v
            land = [x.strip() for x in str(static.get("land area", "")).splitlines() if x.strip()]
            if comps and len(land) == len(comps):
                for comp, lab in zip(comps, land):
                    comp["group"] = lab
            records.append({"project": project, "date": d.isoformat(), "raw": text, "components": comps,
                            "mappedId": ids.get(project, ""), "static": static,
                            "source": "Excel baseline", "sourceCell": f"Sheet1!{cell.coordinate}", "corrections": []})

    # a column whose header differs between blocks means someone typed a date by hand
    for c, vals in header_dates.items():
        dates = {d for d, _ in vals}
        if len(dates) > 1:
            col = ws.cell(1, c).column_letter
            detail = "; ".join(f"{d} in {', '.join(sorted(p for dd, p in vals if dd == d))}" for d in sorted(dates))
            issue("warn", "HEADER_DATE_CONFLICT", "", "", f"column {col}", f"Same column carries different dates: {detail}")

    # ---- apply explicit corrections -----------------------------------------------------------
    for fix in cfg.get("corrections", []):
        if not fix.get("apply"):
            continue
        hits = [r for r in records if r["project"] == fix["project"] and r["date"] == fix["date"]]
        if not hits:
            issue("warn", "CORRECTION_NOT_MATCHED", fix["project"], fix["date"], "", f"Correction {fix['id']} matched no record")
            continue
        rec = hits[0]
        comp = rec["components"][fix.get("component", 0)] if rec["components"] else None
        if "set_total" in fix and comp:
            comp["total"] = fix["set_total"]
        if "set_units" in fix and comp:
            comp["units"] = fix["set_units"]
        if comp and ("set_total" in fix or "set_units" in fix):
            comp["raw"] = f"{comp['reportedPct']:g}% {comp['label']} ({comp['units']}/{comp['total']} units)"
        if "set_date" in fix:
            rec["date"] = fix["set_date"]
        rec["corrections"].append(fix["id"])
        issue("info", "CORRECTION_APPLIED", fix["project"], rec["date"], rec["sourceCell"], f"{fix['id']}: {fix['reason']}")

    # ---- group labels where the Land Area row does not give one layout per group ----------------
    for rec in records:
        labels = cfg.get("group_labels", {}).get(rec["project"])
        if labels and len(labels) == len(rec["components"]):
            for comp, lab in zip(rec["components"], labels):
                comp["group"] = lab

    # ---- derived fields, same shape as the collector's records ---------------------------------
    for rec in records:
        sg = [c for c in rec["components"] if c["label"] == "Signed"]
        bk = [c for c in rec["components"] if c["label"] == "Booked"]
        comps = rec["components"]
        rec["signed"] = sum(c["units"] for c in sg) if comps else None
        rec["booked"] = sum(c["units"] for c in bk) if comps else None
        rec["total"] = sum(c["total"] for c in comps) if comps else None
        rec["signedPct"] = (rec["signed"] / rec["total"] * 100) if comps and rec["total"] else None
    records.sort(key=lambda r: (r["project"], r["date"]))

    # ---- validation on the corrected data --------------------------------------------------------
    seen = set()
    for r in records:
        k = (r["project"], r["date"])
        if k in seen:
            issue("error", "DUPLICATE_DATE", r["project"], r["date"], r["sourceCell"], "Two records for the same project and date")
        seen.add(k)
        for i, c in enumerate(r["components"]):
            calc = c["units"] / c["total"] * 100 if c["total"] else 0
            if c["units"] > c["total"]:
                issue("error", "UNITS_EXCEED_TOTAL", r["project"], r["date"], r["sourceCell"], f"group {i+1}: {c['units']} > {c['total']}")
            elif abs(calc - c["reportedPct"]) > PCT_TOLERANCE:
                issue("warn", "PCT_MISMATCH", r["project"], r["date"], r["sourceCell"],
                      f"group {i+1}: {c['units']}/{c['total']} is {calc:.1f}% but the cell says {c['reportedPct']:g}%")
    by_proj = defaultdict(list)
    for r in records:
        if r["components"]:
            by_proj[r["project"]].append(r)
    for proj, rs in by_proj.items():
        rs.sort(key=lambda r: r["date"])
        for prev, cur in zip(rs, rs[1:]):
            if len(prev["components"]) != len(cur["components"]):
                issue("warn", "GROUP_COUNT_CHANGED", proj, cur["date"], cur["sourceCell"],
                      f"{len(prev['components'])} groups on {prev['date']} vs {len(cur['components'])} now")
                continue
            for i, (a, b) in enumerate(zip(prev["components"], cur["components"])):
                if a["total"] != b["total"]:
                    issue("info", "TOTAL_CHANGED", proj, cur["date"], cur["sourceCell"], f"group {i+1}: {a['total']} -> {b['total']}; signed percentage is recalculated using the recorded total; denominator change flagged for user review")
                if a["label"] == b["label"] == "Signed" and b["units"] < a["units"]:
                    issue("warn", "SIGNED_DECREASE", proj, cur["date"], cur["sourceCell"],
                          f"group {i+1}: {a['units']} on {prev['date']} -> {b['units']} (cancellations, or a typo)")

    out = {"schemaVersion": 2, "source": Path(xlsx).name, "generatedFrom": "scripts/import_history_from_excel.py", "records": records}
    (DATA / "history-seed.json").write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    with (DATA / "history-issues.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["severity", "code", "project", "date", "cell", "detail"])
        w.writeheader()
        w.writerows(sorted(issues, key=lambda i: ({"error": 0, "warn": 1, "info": 2}[i["severity"]], i["project"], i["date"])))
    n = defaultdict(int)
    for i in issues:
        n[i["severity"]] += 1
    print(f"{len(records)} records from {len(blocks)} projects -> data/history-seed.json")
    print(f"issues: {n['error']} error, {n['warn']} warn, {n['info']} info -> data/history-issues.csv")
    sys.exit(1 if n["error"] else 0)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
