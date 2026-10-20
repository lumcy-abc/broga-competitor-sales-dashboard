#!/usr/bin/env python3
"""Generate long-term monthly daily trackers and monthly weekly reports.

Run after scripts/export_competitor_workbook.mjs. Reads the generated workbook's
Historical Records sheet; does not alter the TEDUH collector or its JSON history.
"""
from datetime import datetime, date, timedelta
from pathlib import Path
from collections import defaultdict
from copy import copy
from calendar import monthrange
from zoneinfo import ZoneInfo
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path.cwd()
SOURCE = ROOT / "Broga_Competitor_Sales_Analysis.xlsx"
DAILY_OUT = ROOT / "Broga_Daily_Tracking.xlsx"
WEEKLY_OUT = ROOT / "Broga_Weekly_Tracking.xlsx"
GREEN, BROWN, SAGE, CREAM = "526747", "8B6F4E", "E8EDE3", "F3EBDD"
WHITE, DARK, GRAY, LINE, PALE, AMBER, LIGHT = "FFFFFF", "263126", "666666", "D7DCCF", "DCE8D5", "F6E5C9", "F7F8F4"
TODAY = datetime.now(ZoneInfo("Asia/Kuala_Lumpur")).date()


def read_records():
    if not SOURCE.exists():
        raise FileNotFoundError(f"Missing {SOURCE.name}. Run the existing Excel export first.")
    book = load_workbook(SOURCE, read_only=True, data_only=True)
    if "Historical Records" not in book.sheetnames:
        raise ValueError("Source workbook has no Historical Records sheet.")
    sheet = book["Historical Records"]
    headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    out = {}
    skipped_dates = skipped_status = skipped_fields = duplicates = 0
    for row_num, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        rec = dict(zip(headers, values))
        raw_date = rec.get("Record Date")
        try:
            if isinstance(raw_date, datetime):
                dt = raw_date.date()
            elif isinstance(raw_date, date):
                dt = raw_date
            else:
                raw = str(raw_date).strip()
                try:
                    dt = datetime.strptime(raw, "%d/%m/%Y").date()
                except ValueError:
                    dt = datetime.strptime(raw[:10], "%Y-%m-%d").date()
        except (ValueError, TypeError, AttributeError):
            skipped_dates += 1
            continue
        status_type = str(rec.get("Status Type", "")).strip().lower()
        if status_type != "signed":
            if status_type not in ("booked", "yet to launch", "yet to sign", "unclassified", ""):
                print(f"WARNING: row {row_num}: unknown status type {status_type!r}; excluded from signed tracker.")
            skipped_status += 1
            continue
        try:
            signed, total = float(rec.get("Signed Units")), float(rec.get("Total Units"))
        except (ValueError, TypeError):
            skipped_fields += 1
            print(f"WARNING: row {row_num}: signed record has invalid unit counts; skipped.")
            continue
        if total <= 0 or signed < 0 or signed > total:
            skipped_fields += 1
            print(f"WARNING: row {row_num}: invalid counts signed={signed}, total={total}; skipped.")
            continue
        project = str(rec.get("Project") or "").strip()
        group = str(rec.get("Unit Group / Layout") or "").strip()
        if not project or not group:
            skipped_fields += 1
            print(f"WARNING: row {row_num}: missing project/layout; skipped.")
            continue
        key = (dt, project, group)
        candidate = {"date": dt, "project": project, "group": group, "signed": signed, "total": total}
        if key in out:
            duplicates += 1
            if out[key]["signed"] == signed and out[key]["total"] == total:
                print(f"WARNING: identical duplicate dated group row {key!r}; duplicate ignored (source row {row_num}).")
                continue
            raise ValueError(
                f"Conflicting duplicate dated group row {key!r}: existing signed/total="
                f"{out[key]['signed']}/{out[key]['total']}, new={signed}/{total} at source row {row_num}. "
                "Resolve the source conflict instead of silently choosing one value."
            )
        out[key] = candidate
    print(f"Source validation: {len(out)} unique signed rows; skipped {skipped_dates} invalid dates, {skipped_status} non-signed/unknown statuses, {skipped_fields} invalid signed rows; {duplicates} duplicate keys.")
    return sorted(out.values(), key=lambda r: (r["date"], r["project"], r["group"]))


def read_history_rows():
    """Read every dated row in Historical Records, including non-signed/unknown statuses.

    Trend calculations still use only valid signed observations, but Snapshot History
    must expose all available source dates/statuses for auditing and date filtering.
    """
    book = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = book["Historical Records"]
    headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row_num, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        rec = dict(zip(headers, values))
        raw_date = rec.get("Record Date")
        try:
            if isinstance(raw_date, datetime): dt = raw_date.date()
            elif isinstance(raw_date, date): dt = raw_date
            else:
                raw = str(raw_date).strip()
                try: dt = datetime.strptime(raw, "%d/%m/%Y").date()
                except ValueError: dt = datetime.strptime(raw[:10], "%Y-%m-%d").date()
        except (ValueError, TypeError, AttributeError):
            print(f"WARNING: history row {row_num} has an invalid date and cannot appear in date filtering: {raw_date!r}")
            continue
        project = str(rec.get("Project") or "").strip()
        group = str(rec.get("Unit Group / Layout") or "").strip()
        status_type = str(rec.get("Status Type") or "Unclassified").strip()
        if not project or not group:
            print(f"WARNING: history row {row_num} has no project/layout label; retained with a placeholder.")
            project = project or "Unidentified project"
            group = group or "Unidentified unit group"
        try: signed = float(rec.get("Signed Units"))
        except (ValueError, TypeError): signed = None
        try: total = float(rec.get("Total Units"))
        except (ValueError, TypeError): total = None
        rows.append({"date": dt, "project": project, "group": group, "status": status_type,
                     "signed": signed, "total": total,
                     "raw": str(rec.get("Raw Recorded Status") or ""),
                     "source": str(rec.get("Source") or "")})
        if status_type.lower() in ("", "unclassified"):
            print(f"WARNING: history row {row_num} has a blank/unclassified status ({project} / {group}, {dt}); retained for review.")
    book.close()
    return rows


def style_cell(cell, fill=WHITE, bold=False, color=DARK, horizontal="center", size=9):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
    thin = Side(style="thin", color=LINE)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def title(ws, last_col, heading, subtitle):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c = ws.cell(1, 1, heading); c.fill = PatternFill("solid", fgColor=GREEN)
    c.font = Font(name="Aptos Display", size=15, bold=True, color=WHITE)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c = ws.cell(2, 1, subtitle); c.fill = PatternFill("solid", fgColor=CREAM)
    c.font = Font(italic=True, color="8B5E34", size=10)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 27
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def status(r):
    return f"{round(r['signed'] / r['total'] * 100)}% Signed ({int(r['signed'])}/{int(r['total'])} units)"


def delta(cur, prev):
    """Describe the change against the previous valid observation, including its gap."""
    if not cur: return "No record"
    if not prev: return "No prior record"
    gap = (cur["date"] - prev["date"]).days
    gap_label = f"over {gap} day(s)" if gap > 1 else "over 1 day"
    if cur["total"] != prev["total"]:
        movement = "Total increased" if cur["total"] > prev["total"] else "Total decreased"
        return f"{movement} — review ({prev['total']} → {cur['total']} units; {prev['date']:%d %b} → {cur['date']:%d %b}; {gap_label})"
    n = int(cur["signed"] - prev["signed"])
    change = "No change" if n == 0 else f"{n:+d} signed"
    return f"{change} ({prev['date']:%d %b} → {cur['date']:%d %b}; {gap_label})"


def daily_change_and_gap(cur, prev):
    if not cur:
        return "No observation", None
    if not prev:
        return "No prior observation", None
    gap = (cur["date"] - prev["date"]).days
    if cur["total"] != prev["total"]:
        movement = "Total increased" if cur["total"] > prev["total"] else "Total decreased"
        return f"{movement} — review scope/release ({int(prev['total'])} → {int(cur['total'])} units)", gap
    change = int(cur["signed"] - prev["signed"])
    return ("No change" if change == 0 else f"{change:+d} signed"), gap


def month_range(start, end):
    y, m = start.year, start.month
    while (y, m) <= (end.year, end.month):
        yield y, m
        if m == 12: y, m = y + 1, 1
        else: m += 1


def month_name(y, m):
    return date(y, m, 1).strftime("%b %Y")


def group_records(records):
    groups = sorted({(r["project"], r["group"]) for r in records}, key=lambda x: (x[0].casefold(), x[1].casefold()))
    lookup = {(r["date"], (r["project"], r["group"])): r for r in records}
    by_group = defaultdict(list)
    for r in records: by_group[(r["project"], r["group"])].append(r)
    for g in by_group: by_group[g].sort(key=lambda r: r["date"])
    return groups, lookup, by_group


def write_group_headers(ws, groups, first_header_row=4):
    # Column A is explicitly labelled PROJECT NAMES per the requested design.
    ws.cell(first_header_row, 1, "PROJECT NAMES")
    ws.cell(first_header_row + 1, 1, "Unit group / layout")
    col = 2
    while col <= len(groups) + 1:
        project = groups[col - 2][0]
        end = col + 1
        while end <= len(groups) + 1 and groups[end - 2][0] == project:
            end += 1
        if end - 1 > col:
            ws.merge_cells(start_row=first_header_row, start_column=col, end_row=first_header_row, end_column=end - 1)
        ws.cell(first_header_row, col, project)
        for c in range(col, end):
            ws.cell(first_header_row + 1, c, groups[c - 2][1])
            ws.column_dimensions[get_column_letter(c)].width = 23
        col = end
    for c in range(1, len(groups) + 2):
        style_cell(ws.cell(first_header_row, c), GREEN, True, WHITE, size=9)
        style_cell(ws.cell(first_header_row + 1, c), BROWN, True, WHITE, size=9)
    ws.column_dimensions["A"].width = 24
    ws.row_dimensions[first_header_row].height = 35
    ws.row_dimensions[first_header_row + 1].height = 32


def build_previous_lookup(by_group, groups, month_start, month_end):
    # For each calendar date, cache the latest observation before that date.
    result = {}
    for g in groups:
        rows = by_group.get(g, [])
        idx = 0
        previous = None
        dt = month_start
        while dt <= month_end:
            while idx < len(rows) and rows[idx]["date"] < dt:
                previous = rows[idx]
                idx += 1
            result[(dt, g)] = previous
            dt += timedelta(days=1)
    return result



def outline_project_blocks(ws, groups, first_data_row, last_data_row, header_row=5):
    """Add a clear vertical boundary around each project's layout columns."""
    medium = Side(style="medium", color=GREEN)
    for idx, group in enumerate(groups, start=2):
        starts_project = idx == 2 or groups[idx - 3][0] != group[0]
        ends_project = idx == len(groups) + 1 or groups[idx - 1][0] != group[0]
        if starts_project or ends_project:
            col = idx
            for row in range(header_row, last_data_row + 1):
                cell = ws.cell(row, col)
                b = cell.border
                cell.border = Border(left=medium if starts_project else b.left,
                                     right=medium if ends_project else b.right,
                                     top=b.top, bottom=b.bottom)


def create_daily(records, groups, lookup, by_group):
    wb = Workbook()
    wb.remove(wb.active)
    # Daily tabs are created only for months with valid observations, plus the current month.
    # This preserves useful history without producing empty tabs for months with no records.
    data_months = {(r["date"].year, r["date"].month) for r in records}
    data_months.add((TODAY.year, TODAY.month))
    months = sorted(data_months)
    for y, m in months:
        ws = wb.create_sheet(month_name(y, m))
        last_day = TODAY.day if (y, m) == (TODAY.year, TODAY.month) else monthrange(y, m)[1] if date(y, m, 1) < TODAY.replace(day=1) else 0
        month_start = date(y, m, 1)
        month_end = date(y, m, last_day) if last_day else date(y, m, monthrange(y, m)[1])
        title(ws, len(groups) + 1, f"BROGA COMPETITOR SALES ANALYSIS | DAILY TRACKER | {month_name(y,m).upper()}",
              "Daily validation view · Sales status is combined with unit count and percentage. Daily Change compares with the previous available signed observation; it does not sum snapshots.")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(groups)+1)
        ws.cell(3, 1, "A new monthly tab is generated automatically when this report runs in a new month. Historical month tabs are retained in the regenerated workbook.")
        ws.cell(3, 1).font = Font(italic=True, size=9, color=DARK); ws.cell(3,1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[3].height = 29
        write_group_headers(ws, groups, 5)
        previous_lookup = build_previous_lookup(by_group, groups, month_start, month_end)
        rr = 7
        for day_num in range(1, month_end.day + 1):
            dt = date(y, m, day_num)
            ws.cell(rr, 1, dt.strftime("%d %b %Y") + " · Sales Status")
            ws.cell(rr + 1, 1, "Change vs Previous Observation")
            ws.cell(rr + 2, 1, "Days Since Previous Observation")
            style_cell(ws.cell(rr, 1), BROWN, True, WHITE, "left")
            style_cell(ws.cell(rr + 1, 1), CREAM, True, DARK, "left")
            style_cell(ws.cell(rr + 2, 1), SAGE, True, DARK, "left")
            for j, g in enumerate(groups, 2):
                cur = lookup.get((dt, g))
                prev = previous_lookup.get((dt, g))
                val = status(cur) if cur else ("No signed observation" if dt <= TODAY else "Not yet reached")
                change, gap = daily_change_and_gap(cur, prev)
                style_cell(ws.cell(rr, j, val), SAGE if cur else LIGHT)
                if not cur:
                    change = "No observation" if dt <= TODAY else "Not yet reached"
                if "review" in change.lower():
                    fill = AMBER
                elif change.startswith("+"):
                    fill = PALE
                elif change.startswith("-"):
                    fill = "FCE4D6"
                else:
                    fill = LIGHT
                style_cell(ws.cell(rr + 1, j, change), fill, True)
                style_cell(ws.cell(rr + 2, j, gap if gap is not None else "—"), LIGHT)
            ws.row_dimensions[rr].height = 32; ws.row_dimensions[rr+1].height = 27; ws.row_dimensions[rr+2].height = 22
            rr += 3
        ws.freeze_panes = "B7"
        ws.print_title_rows = "1:6"
        ws.auto_filter.ref = f"A6:{get_column_letter(len(groups)+1)}{rr-1}"
        outline_project_blocks(ws, groups, 7, rr - 1, header_row=5)
    wb.active = len(wb.sheetnames) - 1
    wb.save(DAILY_OUT)


def month_week_periods(y, m):
    first = date(y, m, 1)
    last = date(y, m, monthrange(y, m)[1])
    # Include every Monday–Sunday week that overlaps this month, even when the
    # Monday falls in the preceding month. The same boundary week may appear in both tabs.
    monday = first - timedelta(days=first.weekday())
    periods = []
    start = monday
    while start <= last:
        end = start + timedelta(days=6)
        periods.append((start, end))
        start += timedelta(days=7)
    return periods


def create_weekly(records, groups, lookup, by_group):
    wb = Workbook(); wb.remove(wb.active)
    # Weekly comparisons begin in October 2026. Create a monthly tab only when
    # that month has valid records; future tabs appear automatically after data arrives.
    weekly_start = date(2026, 10, 1)
    months = sorted({(r["date"].year, r["date"].month) for r in records
                     if r["date"] >= weekly_start})
    if not months:
        raise ValueError("No valid signed observations on or after 1 October 2026; weekly report has not been created.")
    for y, m in months:
        ws = wb.create_sheet(month_name(y, m))
        # Weekly tracker uses a compact one-row title, matching the daily tracker's
        # clear visual hierarchy while removing the two redundant explanatory rows.
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(groups)+1)
        c = ws.cell(1, 1, f"BROGA COMPETITOR SALES ANALYSIS | WEEKLY TRACKER | {month_name(y,m).upper()}")
        c.fill = PatternFill("solid", fgColor=GREEN)
        c.font = Font(name="Aptos Display", size=15, bold=True, color=WHITE)
        c.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 32
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        write_group_headers(ws, groups, 2)
        rr=4
        periods=month_week_periods(y,m)
        for wi,(start,end) in enumerate(periods):
            period_end=min(end,TODAY)
            ws.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=len(groups)+1)
            c=ws.cell(rr,1,f"WEEK · {start.day} {start:%b %Y} – {end.day} {end:%b %Y}")
            c.fill=PatternFill("solid",fgColor=BROWN); c.font=Font(bold=True,color=WHITE,size=10); c.alignment=Alignment(vertical="center")
            ws.row_dimensions[rr].height=24; rr+=1
            for metric,kind in [("Sales Status","status"),("Weekly Change","change")]:
                style_cell(ws.cell(rr,1,metric), SAGE if kind=="status" else CREAM, True, DARK, "left")
                for j,g in enumerate(groups,2):
                    valid=[r for r in by_group.get(g,[]) if start<=r["date"]<=period_end] if start<=TODAY else []
                    cur=max(valid,key=lambda x:x["date"]) if valid else None
                    prev_start=start-timedelta(days=7); prev_end=start-timedelta(days=1)
                    prior=[r for r in by_group.get(g,[]) if prev_start<=r["date"]<=prev_end]
                    prev=max(prior,key=lambda x:x["date"]) if prior else None
                    if kind=="status":
                        val=status(cur) if cur else ("In progress" if start<=TODAY<=end else "Future week" if start>TODAY else "No record this week")
                        fill=SAGE if cur else CREAM
                    else:
                        if start>TODAY: val="Future week"
                        elif not cur: val="In progress" if start<=TODAY<=end else "No record this week"
                        elif not prev: val="Baseline / no prior-week record"
                        else: val=delta(cur,prev)
                        fill=PALE if (str(val).startswith("+") or val=="No change") else AMBER if ("review" in str(val).lower() or "total increased" in str(val).lower()) else CREAM
                    style_cell(ws.cell(rr,j,val),fill,kind=="change")
                ws.row_dimensions[rr].height=32 if kind=="status" else 28; rr+=1
        ws.merge_cells(start_row=rr+1,start_column=1,end_row=rr+1,end_column=len(groups)+1)
        c=ws.cell(rr+1,1,"Monthly tabs are generated from available history. Weeks are Monday–Sunday; weeks spanning month boundaries are shown in each month they overlap.")
        c.fill=PatternFill("solid",fgColor=CREAM); c.font=Font(italic=True,color="8B5E34",size=9); c.alignment=Alignment(wrap_text=True,vertical="center")
        ws.row_dimensions[rr+1].height=30
        ws.freeze_panes="B4"; ws.print_title_rows="1:3"
        outline_project_blocks(ws, groups, 4, rr + 1, header_row=2)
    wb.active=len(wb.sheetnames)-1
    wb.save(WEEKLY_OUT)



def create_snapshot_history(records, by_group, history_rows):
    """Create a date-filterable history including all source statuses and valid signed trends."""
    wb = load_workbook(WEEKLY_OUT)
    sheet_name = "Snapshot History"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    if "Latest Snapshot Comparison" in wb.sheetnames:
        del wb["Latest Snapshot Comparison"]
    ws = wb.create_sheet(sheet_name, 0)
    headers = [
        "Observation Date", "Project", "Unit Group / Layout", "Status Type", "Signed Units",
        "Total Units", "Signed %", "Previous Observation Date", "Days Since Previous Observation",
        "Change in Signed Units", "Change in Signed Percentage Points", "Source Status Text",
        "Data Quality / Interpretation"
    ]
    title(ws, len(headers), "BROGA COMPETITOR SALES ANALYSIS | SNAPSHOT HISTORY",
          "All dated source-history rows are retained, including non-signed and unclassified statuses. Only valid signed observations contribute to change calculations. Use the filter on Observation Date; other columns intentionally have no filter dropdown.")
    header_row = 4
    for col, header in enumerate(headers, 1):
        style_cell(ws.cell(header_row, col, header), BROWN, True, WHITE, size=9)
    ws.row_dimensions[header_row].height = 42
    rr = header_row + 1
    written_keys = set()
    # Valid signed observations first; retain previous-valid-observation deltas by group.
    for group_key in sorted(by_group, key=lambda x: (x[0].casefold(), x[1].casefold())):
        observations = sorted(by_group[group_key], key=lambda x: x["date"])
        previous = None
        for current in observations:
            key = (current["date"], current["project"], current["group"])
            written_keys.add(key)
            gap = (current["date"] - previous["date"]).days if previous else None
            signed_change = current["signed"] - previous["signed"] if previous else None
            current_pct = current["signed"] / current["total"] if current["total"] else None
            previous_pct = previous["signed"] / previous["total"] if previous and previous["total"] else None
            pp_change = (current_pct - previous_pct) * 100 if current_pct is not None and previous_pct is not None else None
            notes = []
            if previous is None:
                notes.append("Baseline: no earlier valid signed observation")
            else:
                if current["total"] != previous["total"]:
                    notes.append("Total-unit denominator changed; review release/scope")
                if gap > 1:
                    notes.append(f"{gap}-day gap; change is not daily movement")
                if signed_change < 0:
                    notes.append("Signed count decreased; verify source/status correction")
            if not notes: notes.append("Comparable signed observation")
            values = [current["date"], current["project"], current["group"], "Signed",
                      current["signed"], current["total"], current_pct,
                      previous["date"] if previous else None, gap, signed_change, pp_change,
                      "", "; ".join(notes)]
            for col, value in enumerate(values, 1):
                cell = ws.cell(rr, col, value)
                style_cell(cell, WHITE if rr % 2 else LIGHT, color=DARK,
                           horizontal="left" if col in (2, 3, 12, 13) else "center", size=9)
            ws.cell(rr, 1).number_format = "dd mmm yyyy"
            ws.cell(rr, 6).number_format = "0"
            ws.cell(rr, 7).number_format = "0.0%"
            ws.cell(rr, 8).number_format = "dd mmm yyyy"
            ws.cell(rr, 10).number_format = "+0;-0;0"
            ws.cell(rr, 11).number_format = "+0.0;-0.0;0.0"
            if signed_change is not None:
                ws.cell(rr, 10).font = Font(name="Aptos", size=9, bold=True, color=("C00000" if signed_change < 0 else DARK))
            if pp_change is not None:
                ws.cell(rr, 11).font = Font(name="Aptos", size=9, bold=True, color=("C00000" if pp_change < 0 else DARK))
            rr += 1
            previous = current
    # Keep dated non-signed, unknown-status, and invalid-signed rows for date completeness/audit.
    for item in sorted(history_rows, key=lambda x: (x["date"], x["project"].casefold(), x["group"].casefold(), x["status"].casefold())):
        key = (item["date"], item["project"], item["group"])
        status_lower = item["status"].strip().lower()
        valid_signed = status_lower == "signed" and item["signed"] is not None and item["total"] is not None and item["total"] > 0 and 0 <= item["signed"] <= item["total"]
        if valid_signed and key in written_keys:
            continue
        notes = []
        if status_lower in ("", "unclassified"):
            notes.append("Blank/unclassified status; verify source")
        elif status_lower != "signed":
            notes.append("Non-signed status retained for audit; excluded from signed trend")
        else:
            notes.append("Signed status has invalid/missing unit counts; verify source")
        pct = item["signed"] / item["total"] if valid_signed else None
        values = [item["date"], item["project"], item["group"], item["status"],
                  item["signed"] if item["signed"] is not None else None,
                  item["total"] if item["total"] is not None else None, pct,
                  None, None, None, None, item["raw"], "; ".join(notes)]
        for col, value in enumerate(values, 1):
            cell = ws.cell(rr, col, value)
            style_cell(cell, AMBER if col == 13 else (WHITE if rr % 2 else LIGHT), color=DARK,
                       horizontal="left" if col in (2, 3, 12, 13) else "center", size=9)
        ws.cell(rr, 1).number_format = "dd mmm yyyy"
        ws.cell(rr, 7).number_format = "0.0%"
        rr += 1
    if rr == header_row + 1:
        ws.cell(rr, 1, "No dated source-history rows found")
    else:
        # Keep the audit table in chronological order while preserving row styling.
        stored_rows = []
        for row_idx in range(header_row + 1, rr):
            cells = [ws.cell(row_idx, col) for col in range(1, len(headers) + 1)]
            values = [cell.value for cell in cells]
            styles = [copy(cell._style) for cell in cells]
            key_date = values[0].date() if isinstance(values[0], datetime) else values[0]
            stored_rows.append(((key_date, str(values[1] or "").casefold(), str(values[2] or "").casefold(), str(values[3] or "").casefold()), values, styles))
        stored_rows.sort(key=lambda item: item[0])
        for row_idx, (_, values, styles) in enumerate(stored_rows, start=header_row + 1):
            for col, (value, style) in enumerate(zip(values, styles), start=1):
                cell = ws.cell(row_idx, col, value)
                cell._style = style
    # Only Observation Date gets a filter dropdown.
    ws.auto_filter.ref = f"A{header_row}:A{max(header_row, rr - 1)}"
    ws.freeze_panes = "A5"
    widths = [18, 28, 34, 17, 14, 13, 12, 20, 23, 17, 24, 45, 58]
    for col, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(col)].width = width
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    ws.row_dimensions[2].height = 34
    wb.save(WEEKLY_OUT)

def main():
    records=read_records()
    if not records: raise ValueError("No valid signed records found in Historical Records.")
    history_rows=read_history_rows()
    groups, lookup, by_group = group_records(records)
    create_daily(records, groups, lookup, by_group)
    create_weekly(records, groups, lookup, by_group)
    create_snapshot_history(records, by_group, history_rows)
    daily_tabs = load_workbook(DAILY_OUT, read_only=True).sheetnames
    weekly_tabs = load_workbook(WEEKLY_OUT, read_only=True).sheetnames
    print(f"Created {DAILY_OUT.name} with {len(daily_tabs)} tabs ({', '.join(daily_tabs)})")
    print(f"Created {WEEKLY_OUT.name} with {len(weekly_tabs)} tabs ({', '.join(weekly_tabs)}), including date-filtered snapshot history and Monday–Sunday monthly tabs")
    print(f"Source: {len(records)} valid signed observations across {len(groups)} project/layout groups.")

if __name__ == "__main__":
    main()
